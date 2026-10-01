"""
exporter.py — Enhanced Excel export with multiple sheets and formatting.
Inclui coluna Tipo NF e sheet de Lançamentos Manuais.
"""
import pandas as pd
import io

try:
    from openpyxl.styles import PatternFill, Font, Alignment, numbers
    from openpyxl.utils import get_column_letter
    _OPENPYXL_FULL = True
except ImportError:
    _OPENPYXL_FULL = False

# Cores para formatação Excel
_FILL_HEADER   = "1F3864"   # azul escuro
_FILL_RETORNO  = "FFF2CC"   # amarelo claro — Retorno/Devolução
_FILL_ENTRADA  = "E2EFDA"   # verde claro — Entrada
_FILL_SAIDA    = "FCE4D6"   # laranja claro — Saída
_FONT_WHITE    = "FFFFFF"


def _apply_excel_style(ws):
    """Aplica estilo básico: cabeçalho escuro, fonte branca, auto-largura."""
    if not _OPENPYXL_FULL:
        return
    header_fill = PatternFill("solid", fgColor=_FILL_HEADER)
    header_font = Font(color=_FONT_WHITE, bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    # Auto-fit largura (max 60)
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 60)


def _fmt_currency_col(ws, col_letter, start_row=2):
    """Formata coluna como moeda BRL."""
    if not _OPENPYXL_FULL:
        return
    fmt = '#,##0.00'
    for cell in ws[col_letter][start_row - 1:]:
        if cell.row >= start_row:
            cell.number_format = fmt


def _color_tipo_nf(ws, tipo_col_idx, start_row=2):
    """Coloriza linhas da coluna Tipo NF."""
    if not _OPENPYXL_FULL:
        return
    fills = {
        "Retorno/Devolução": PatternFill("solid", fgColor=_FILL_RETORNO),
        "Entrada":           PatternFill("solid", fgColor=_FILL_ENTRADA),
        "Saída":             PatternFill("solid", fgColor=_FILL_SAIDA),
    }
    for row in ws.iter_rows(min_row=start_row):
        cell = row[tipo_col_idx - 1]
        fill = fills.get(str(cell.value or ""))
        if fill:
            cell.fill = fill


def export_to_excel(invoices, categories=None, obras=None, approvals=None,
                    lancamentos_manuais=None, **kwargs):
    """
    Exporta dados para Excel multi-planilha:
      Notas (Parcelas)     — uma linha por parcela de NF
      Resumo por Nota      — uma linha por NF
      Itens Detalhados     — um item por linha
      Resumo por Obra      — agrupado
      Resumo por Categoria — agrupado
      Lançamentos Manuais  — lançamentos inseridos manualmente (mesmo formato)
    """
    if not invoices and not lancamentos_manuais:
        return None

    categories         = categories or {}
    obras              = obras or {}
    approvals          = approvals or {}
    codigos            = kwargs.get("codigos", {})
    lancamentos_manuais = lancamentos_manuais or []

    summary_rows = []
    invoice_rows = []

    for d in invoices:
        uid     = d.get("uid", d.get("nNF", ""))
        cat     = categories.get(uid, d.get("categoria", "Outros"))
        obra    = obras.get(uid, d.get("obra_sugerida", "Geral"))
        codigo  = codigos.get(uid, "")
        tipo_nf = d.get("tipo_nf_classificado", d.get("tipo", "Entrada"))

        # Remove prefixo "~" do fallback para exibição limpa no Excel
        codigo_limpo = codigo.lstrip("~") if codigo else ""

        invoice_rows.append({
            "Tipo NF":            tipo_nf,
            "Fornecedor":         d.get("fornecedor", ""),
            "NF":                 d.get("nNF", ""),
            "Série":              d.get("serie", ""),
            "Data Emissão NF":    d.get("data_emissao", ""),
            "Valor Total NF":     d.get("valor_total", 0),
            "CNPJ Fornecedor":    d.get("fornecedor_cnpj", ""),
            "Cliente":            d.get("cliente", ""),
            "Codigo Contabil":    codigo_limpo,
            "Categoria":          cat,
            "Obra":               obra,
            "CNO":                d.get("cno_sugerido", ""),
            "Status":             approvals.get(uid, "Pendente"),
            "Chave Acesso":       d.get("chave", ""),
        })

        parcelas = d.get("parcelas", [])
        if not parcelas:
            parcelas = [{"nDup": "1", "dVenc": d.get("data_emissao"), "vDup": d.get("valor_total")}]

        for p in parcelas:
            summary_rows.append({
                "Tipo NF":               tipo_nf,
                "Fornecedor":            d.get("fornecedor", ""),
                "Vencimento Parcela":    p.get("dVenc", ""),
                "Valor Parcela":         float(p.get("vDup") or d.get("valor_total") or 0.0),
                "NF":                    d.get("nNF", ""),
                "Parcela":               p.get("nDup", "1"),
                "Série":                 d.get("serie", ""),
                "Data Emissão NF":       d.get("data_emissao", ""),
                "CNPJ Fornecedor":       d.get("fornecedor_cnpj", ""),
                "UF Fornecedor":         d.get("fornecedor_uf", ""),
                "Cliente":               d.get("cliente", ""),
                "Doc Cliente":           d.get("cliente_doc", ""),
                "Natureza Operação":     d.get("natureza_op", ""),
                "Valor Produtos":        d.get("valor_produtos", 0),
                "Valor Frete":           d.get("valor_frete", 0),
                "Valor Desconto":        d.get("valor_desconto", 0),
                "Valor Total NF":        d.get("valor_total", 0),
                "Codigo Contabil":       codigo_limpo,
                "Forma Pagamento Base":  d.get("pagamento", ""),
                "Categoria":             cat,
                "Obra":                  obra,
                "CNO":                   d.get("cno_sugerido", ""),
                "Status":                approvals.get(uid, "Pendente"),
                "Qtd Itens":             d.get("qtd_itens", 0),
                "Info Adicional":        (d.get("infAdic", "") or "")[:500],
                "Chave Acesso":          d.get("chave", ""),
                "Protocolo":             d.get("protocolo", ""),
            })

    df_summary  = pd.DataFrame(summary_rows)
    df_invoices = pd.DataFrame(invoice_rows)

    # ── Itens Detalhados ────────────────────────────────────────
    detail_rows = []
    for d in invoices:
        uid     = d.get("uid", d.get("nNF", ""))
        tipo_nf = d.get("tipo_nf_classificado", d.get("tipo", "Entrada"))
        codigo  = codigos.get(uid, "").lstrip("~")
        for item in d.get("itens", []):
            detail_rows.append({
                "Tipo NF":        tipo_nf,
                "NF":             d.get("nNF", ""),
                "Fornecedor":     d.get("fornecedor", ""),
                "Obra":           obras.get(uid, d.get("obra_sugerida", "Geral")),
                "Categoria NF":   categories.get(uid, d.get("categoria", "Outros")),
                "Codigo Contabil": codigo,
                "CNO":            d.get("cno_sugerido", ""),
                "Código Produto": item.get("codigo", ""),
                "Descrição":      item.get("descricao", ""),
                "NCM":            item.get("ncm", ""),
                "CFOP":           item.get("cfop", ""),
                "Quantidade":     item.get("qtd", ""),
                "Unidade":        item.get("unidade", ""),
                "Valor Unitário": item.get("valor_unit", ""),
                "Valor Total":    item.get("valor_total", ""),
            })
    df_detail = pd.DataFrame(detail_rows) if detail_rows else pd.DataFrame()

    # ── Resumo por Obra ─────────────────────────────────────────
    if not df_summary.empty:
        df_obra = df_summary.groupby("Obra").agg(
            Quantidade_NFs=("NF", "nunique"),
            Valor_Total=("Valor Parcela", "sum"),
        ).reset_index()
    else:
        df_obra = pd.DataFrame()

    # ── Resumo por Categoria ─────────────────────────────────────
    if not df_summary.empty:
        df_cat = df_summary.groupby("Categoria").agg(
            Quantidade_NFs=("NF", "nunique"),
            Valor_Total=("Valor Parcela", "sum"),
        ).reset_index()
    else:
        df_cat = pd.DataFrame()

    # ── Lançamentos Manuais ──────────────────────────────────────
    manual_rows = []
    for m in lancamentos_manuais:
        manual_rows.append({
            "Tipo NF":            "Manual",
            "Fornecedor":         m.get("fornecedor", ""),
            "Vencimento Parcela": m.get("vencimento", ""),
            "Valor Parcela":      float(m.get("valor", 0) or 0),
            "NF":                 m.get("nr_documento", ""),
            "Parcela":            "1",
            "Série":              "",
            "Data Emissão NF":    m.get("data", ""),
            "CNPJ Fornecedor":    m.get("cnpj", ""),
            "UF Fornecedor":      "",
            "Cliente":            m.get("cliente", ""),
            "Doc Cliente":        "",
            "Natureza Operação":  "LANÇAMENTO MANUAL",
            "Valor Produtos":     float(m.get("valor", 0) or 0),
            "Valor Frete":        0,
            "Valor Desconto":     0,
            "Valor Total NF":     float(m.get("valor", 0) or 0),
            "Codigo Contabil":    m.get("codigo_contabil", "").lstrip("~"),
            "Forma Pagamento Base": m.get("forma_pagamento", ""),
            "Categoria":          m.get("categoria", ""),
            "Obra":               m.get("obra", ""),
            "CNO":                m.get("cno", ""),
            "Status":             "Aprovado",
            "Qtd Itens":          0,
            "Info Adicional":     m.get("observacao", ""),
            "Chave Acesso":       "",
            "Protocolo":          "",
        })
    df_manual = pd.DataFrame(manual_rows) if manual_rows else pd.DataFrame()

    # ── Write to Excel ──────────────────────────────────────────
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df_summary.to_excel(writer, index=False, sheet_name="Notas (Parcelas)")
        df_invoices.to_excel(writer, index=False, sheet_name="Resumo por Nota")
        if not df_detail.empty:
            df_detail.to_excel(writer, index=False, sheet_name="Itens Detalhados")
        if not df_obra.empty:
            df_obra.to_excel(writer, index=False, sheet_name="Resumo por Obra")
        if not df_cat.empty:
            df_cat.to_excel(writer, index=False, sheet_name="Resumo por Categoria")
        if not df_manual.empty:
            df_manual.to_excel(writer, index=False, sheet_name="Lancamentos Manuais")

        if _OPENPYXL_FULL:
            wb = writer.book

            # Estilo e formatação por planilha
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                _apply_excel_style(ws)

                # Detecta colunas monetárias e aplica formato
                headers = [cell.value for cell in ws[1]]
                currency_cols = [
                    "Valor Parcela", "Valor Total NF", "Valor Produtos",
                    "Valor Frete", "Valor Desconto", "Valor_Total",
                ]
                for col_name in currency_cols:
                    if col_name in headers:
                        col_idx = headers.index(col_name) + 1
                        _fmt_currency_col(ws, get_column_letter(col_idx))

                # Coloriza coluna Tipo NF
                if "Tipo NF" in headers:
                    tipo_idx = headers.index("Tipo NF") + 1
                    _color_tipo_nf(ws, tipo_idx)

    return output.getvalue()
